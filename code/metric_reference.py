import csv
import math
import os
import re
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd
from PIL import Image

# ========================
# 路径配置
# ========================
TASK_DIR = r"inputs/metric_reference"

# ========================
# 自然头位前后/上下符号约定
# 经 470 例 pre 标注点实测：旋转到自然头位后，解剖前部点 x 恒大于后部点 x（470/470 一致），
# 因此“前方 = +x 方向”；图像坐标系 y 向下增大，故“下方 = +y”。
# ========================
ANTERIOR_X = 1.0  # 前方对应 +x 方向
PRE_MEASUREMENTS_FILE = os.path.join(TASK_DIR, "pre_measurements.txt")
DELTA_MEASUREMENTS_FILE = os.path.join(TASK_DIR, "post-pre_measurements.txt")
POST_MEASUREMENTS_FILE = os.path.join(TASK_DIR, "post_measurements.txt")
PATIENT_INFO_EXCEL = os.path.join(TASK_DIR, "指标预测指标-患者基本信息.xlsx")

PRE_POINTS_CSV = os.path.join(TASK_DIR, "pre.csv")
POST_POINTS_CSV = os.path.join(TASK_DIR, "post.csv")
ROTATION_CSV = os.path.join(TASK_DIR, "rotation.csv")
OUTPUT_DIR = os.path.join("outputs", "metric_reference")
IMAGE_DIR = os.path.join("inputs", "images")

# ========================
# 可视化配置
# ========================
# 构建数据集时，把计算所用的旋转后标志点可视化保存，便于核对数据
ENABLE_VISUALIZATION = False
VIS_OUTPUT_DIR = os.path.join(OUTPUT_DIR, "vis_landmark")
VIS_POINT_RADIUS = 8
VIS_POINT_COLOR = (0, 0, 255)   # BGR: 亮红色
VIS_TEXT_COLOR = (0, 255, 255)  # BGR: 亮黄色
VIS_FONT_SCALE = 0.7

_image_center_cache = {}

TRAIN_BATCHES = {"1", "2", "3", "4", "5", "6"}
TEST_BATCHES = {"7"}
IMAGE_EXT = ".bmp"

# Metric order and geometric definitions used by the upstream measurement implementation.
ALL_METRIC_NAMES = [
    "SN-FH(°)", "BaN-FH(°)", "S-N(mm)", "S-Ba(mm)",
    "SNA(°)", "PP-FH(°)", "ANS-N(自然)(mm)", "A-N(自然)(mm)", "A-NP(mm)", "PNS-A(mm)", "Co-A(mm)",
    "Rlc-Go-Me(°)", "Rlc-Go-N(°)", "N-Go-Me(°)", "N-S-Rlc(°)", "S-Rlc-Go(°)", "SNB(°)",
    "Pog-N(自然)(mm)", "Pog-NB(mm)", "Pog-B(自然)(mm)", "B-N(自然)(mm)", "B-NP(mm)", "Pog-NP(mm)",
    "MP-FH(°)", "MP-SN(°)", "N-S-Gn(°)", "SGn-BaN(°)", "SGn-FH(°)", "FH-NPog(°)", "Go-Pog(mm)", "Co-Go(mm)", "Co-Gn(mm)",
    "ANB(°)", "PP-GoGn(°)",
    "S-Go/N-Me(%)", "ANS-Me/N-Me(%)", "N-ANS(mm)",
    "U1-A(mm)", "U1-NA(mm)", "U1-PP(mm)", "U1-APog(mm)",
    "U1p-N(自然)(mm)", "U1p-Pog(自然)(mm)", "U1ta-N(自然)(mm)", "U1ta-Pog(自然)(mm)",
    "U1c-N(自然)(mm)", "U1c-Pog(自然)(mm)", "L1-Pog(mm)", "L1-NB(mm)", "L1-MP(mm)", "L1-APog(mm)",
    "L1p-N(自然)(mm)", "L1p-Pog(自然)(mm)", "L1c-N(自然)(mm)", "L1c-Pog(自然)(mm)", "Holdaway ratio",
    "U1-SN(°)", "U1-NA(°)", "L1-NB(°)", "L1-MP(°)", "L1Pog-OP(°)", "U1-L1(°)",
    "U6-PP(mm)", "L6-MP(mm)",
    "overbite(mm)", "overjet(mm)", "OPa-SN(°)", "OPa-FH(°)", "OPf-SN(°)", "OPf-FH(°)", "L1-OP(mm)",
    "sSn-glabella(自然)(mm)", "UL-Sn(自然)(mm)", "LL-Sn(自然)(mm)", "Ms-Sn(自然)(mm)", "sPog-Sn(自然)(mm)",
    "Sn-Stoms(mm)", "LL-sPogUL(mm)", "Gl-Sn-sPog(°)",
    "ANS-Stoms(mm)", "Sn-A(自然)(mm)", "Sn-N h(自然)(mm)", "Sn-N v(自然)(mm)", "Sn-ANS h(自然)(mm)", "Sn-ANS v(自然)(mm)",
    "U1p-Sn(自然)(mm)", "U1p-sPog(自然)(mm)", "U1p-UL(mm)",
    "U1ta-Sn(自然)(mm)", "U1ta-sPog(自然)(mm)", "U1ta-UL(mm)",
    "U1c-Sn(自然)(mm)", "U1c-sPog(自然)(mm)", "U1c-UL(mm)",
    "L1p-Sn(自然)(mm)", "L1p-sPog(自然)(mm)", "L1p-UL(mm)",
    "L1c-Sn(自然)(mm)", "L1c-sPog(自然)(mm)", "L1c-UL(mm)",
    "Stomi-U1p(mm)", "Stoms-U1p(mm)", "Stomi-L1p(mm)", "Stoms-L1p(mm)",
]


def normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_batch(value) -> str:
    text = normalize_text(value)
    if text == "":
        return ""
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except ValueError:
        pass
    return text


def normalize_metric_name(name: str) -> str:
    text = normalize_text(name)
    text = text.replace("（", "(").replace("）", ")")
    text = text.replace("°", "deg").replace("％", "%")
    text = re.sub(r"\s+", "", text)
    text = text.replace("(自然)", "")
    return text.lower()


CANONICAL_TO_METRIC = {normalize_metric_name(name): name for name in ALL_METRIC_NAMES}

# 针对 txt 中省略单位、空格、全角括号等写法做别名兼容
METRIC_ALIASES = {
    "sn-fh": "SN-FH(°)",
    "sna": "SNA(°)",
    "a-n": "A-N(自然)(mm)",
    "snb": "SNB(°)",
    "pog-n": "Pog-N(自然)(mm)",
    "pog-nb": "Pog-NB(mm)",
    "b-n": "B-N(自然)(mm)",
    "mp-fh": "MP-FH(°)",
    "mp-sn": "MP-SN(°)",
    "n-s-gn": "N-S-Gn(°)",
    "sgn-fh": "SGn-FH(°)",
    "fh-npog": "FH-NPog(°)",
    "anb": "ANB(°)",
    "pp-gogn": "PP-GoGn(°)",
    "s-go/n-me": "S-Go/N-Me(%)",
    "ans-me/n-me": "ANS-Me/N-Me(%)",
    "u1-na": "U1-NA(mm)",
    "u1-apog": "U1-APog(mm)",
    "l1-nb": "L1-NB(mm)",
    "l1-apog": "L1-APog(mm)",
    "u1-sn": "U1-SN(°)",
    "l1-mp": "L1-MP(°)",
    "u6-pp": "U6-PP(mm)",
    "l6-mp": "L6-MP(mm)",
    "overbite": "overbite(mm)",
    "overjet": "overjet(mm)",
    "opa-sn": "OPa-SN(°)",
    "ssn-glabella": "sSn-glabella(自然)(mm)",
    "ul-sn": "UL-Sn(自然)(mm)",
    "ll-sn": "LL-Sn(自然)(mm)",
    "ms-sn": "Ms-Sn(自然)(mm)",
    "spog-sn": "sPog-Sn(自然)(mm)",
    "ll-spogul": "LL-sPogUL(mm)",
    "gl-sn-spog": "Gl-Sn-sPog(°)",
    "sn-a": "Sn-A(自然)(mm)",
    "sn-ansh": "Sn-ANS h(自然)(mm)",
    "sn-ansv": "Sn-ANS v(自然)(mm)",
    "u1p-sn": "U1p-Sn(自然)(mm)",
    "u1p-spog": "U1p-sPog(自然)(mm)",
    "u1ta-sn": "U1ta-Sn(自然)(mm)",
    "u1ta-spog": "U1ta-sPog(自然)(mm)",
    "u1c-sn": "U1c-Sn(自然)(mm)",
    "u1c-spog": "U1c-sPog(自然)(mm)",
    "u1c-ul": "U1c-UL(mm)",
    "l1p-spog": "L1p-sPog(自然)(mm)",
    "l1c-spog": "L1c-sPog(自然)(mm)",
    "l1c-ul": "L1c-UL(mm)",
    "stomi-u1p": "Stomi-U1p(mm)",
    "stoms-u1p": "Stoms-U1p(mm)",
    "stomi-l1p": "Stomi-L1p(mm)",
    "u1-pp": "U1-PP(mm)",
    "u1p-n": "U1p-N(自然)(mm)",
    "u1ta-n": "U1ta-N(自然)(mm)",
    "u1c-n": "U1c-N(自然)(mm)",
}


def resolve_metric_name(raw_name: str) -> str:
    key = normalize_metric_name(raw_name)
    if key in CANONICAL_TO_METRIC:
        return CANONICAL_TO_METRIC[key]
    if key in METRIC_ALIASES:
        return METRIC_ALIASES[key]
    raise KeyError(f"指标清单中的指标无法匹配计算公式: {raw_name}")


def read_measurement_list(path: str) -> List[str]:
    with open(path, "r", encoding="utf-8-sig") as file:
        raw_names = [line.strip() for line in file if line.strip()]
    return [resolve_metric_name(name) for name in raw_names]


# ========================
# 几何函数：参考 landmark_xsx_combinecsv.py，并加入空值保护
# ========================
def is_valid_point(point) -> bool:
    if point is None or len(point) != 2:
        return False
    return not (pd.isna(point[0]) or pd.isna(point[1]))


def safe_div(numerator: float, denominator: float) -> float:
    if denominator is None or denominator == 0 or pd.isna(denominator):
        return np.nan
    return numerator / denominator


def calculate_distance(point1, point2):
    if not is_valid_point(point1) or not is_valid_point(point2):
        return np.nan
    return math.sqrt((point2[0] - point1[0]) ** 2 + (point2[1] - point1[1]) ** 2)


def calculate_foot_to_c_distance(A, B, C):
    if not is_valid_point(A) or not is_valid_point(B) or not is_valid_point(C):
        return np.nan
    Ax, Ay = A
    Bx, By = B
    Cx, Cy = C
    vecBCx = Cx - Bx
    vecBCy = Cy - By
    vecBAx = Ax - Bx
    vecBAy = Ay - By
    dot_product = vecBAx * vecBCx + vecBAy * vecBCy
    bc_length_sq = vecBCx ** 2 + vecBCy ** 2
    if bc_length_sq == 0:
        return math.hypot(Cx - Ax, Cy - Ay)
    t = dot_product / bc_length_sq
    Dx = Bx + t * vecBCx
    Dy = By + t * vecBCy
    return math.hypot(Cx - Dx, Cy - Dy)


def calculate_angle_between_vectors(pointA, pointB, pointC, pointD):
    if not all(is_valid_point(p) for p in [pointA, pointB, pointC, pointD]):
        return np.nan
    vectorAB = [pointB[0] - pointA[0], pointB[1] - pointA[1]]
    vectorCD = [pointD[0] - pointC[0], pointD[1] - pointC[1]]
    dot_product = vectorAB[0] * vectorCD[0] + vectorAB[1] * vectorCD[1]
    magnitudeAB = math.sqrt(vectorAB[0] ** 2 + vectorAB[1] ** 2)
    magnitudeCD = math.sqrt(vectorCD[0] ** 2 + vectorCD[1] ** 2)
    if magnitudeAB == 0 or magnitudeCD == 0:
        return np.nan
    cos_theta = dot_product / (magnitudeAB * magnitudeCD)
    cos_theta = np.clip(cos_theta, -1.0, 1.0)
    return math.degrees(math.acos(cos_theta))


def calculate_angle_at_center(pointA, pointB, pointC):
    if not all(is_valid_point(p) for p in [pointA, pointB, pointC]):
        return np.nan
    vectorBA = [pointA[0] - pointB[0], pointA[1] - pointB[1]]
    vectorBC = [pointC[0] - pointB[0], pointC[1] - pointB[1]]
    dot_product = vectorBA[0] * vectorBC[0] + vectorBA[1] * vectorBC[1]
    magnitudeBA = math.sqrt(vectorBA[0] ** 2 + vectorBA[1] ** 2)
    magnitudeBC = math.sqrt(vectorBC[0] ** 2 + vectorBC[1] ** 2)
    if magnitudeBA == 0 or magnitudeBC == 0:
        return np.nan
    cos_theta = dot_product / (magnitudeBA * magnitudeBC)
    cos_theta = np.clip(cos_theta, -1.0, 1.0)
    return math.degrees(math.acos(cos_theta))


def calculate_natural_pos_level_distance(pointA, pointB):
    if not is_valid_point(pointA) or not is_valid_point(pointB):
        return np.nan
    return pointB[0] - pointA[0]


def calculate_natural_pos_vertical_distance(pointA, pointB):
    if not is_valid_point(pointA) or not is_valid_point(pointB):
        return np.nan
    return pointB[1] - pointA[1]


def calculate_perpendicular_distance_to_line(pointA, pointB, pointC):
    if not all(is_valid_point(p) for p in [pointA, pointB, pointC]):
        return np.nan
    numerator = abs((pointC[0] - pointB[0]) * (pointA[1] - pointB[1]) - (pointA[0] - pointB[0]) * (pointC[1] - pointB[1]))
    denominator = math.sqrt((pointC[0] - pointB[0]) ** 2 + (pointC[1] - pointB[1]) ** 2)
    return safe_div(numerator, denominator)


def calculate_signed_perpendicular_distance(pointA, pointB, pointC):
    if not all(is_valid_point(p) for p in [pointA, pointB, pointC]):
        return np.nan
    xA, yA = pointA
    xB, yB = pointB
    xC, yC = pointC
    if xC == xB:
        yD = yB
    else:
        k = (yC - yB) / (xC - xB)
        yD = k * (xA - xB) + yB
    numerator = abs((xC - xB) * (yA - yB) - (xA - xB) * (yC - yB))
    denominator = math.sqrt((xC - xB) ** 2 + (yC - yB) ** 2)
    distance = safe_div(numerator, denominator)
    return distance if yD > yA else -distance


def calculate_signed_vertical_distance(pointA, pointB, pointC, pointD):
    if not all(is_valid_point(p) for p in [pointA, pointB, pointC, pointD]):
        return np.nan
    bcX = pointC[0] - pointB[0]
    bcY = pointC[1] - pointB[1]
    numerator = bcX * (pointD[0] - pointA[0]) + bcY * (pointD[1] - pointA[1])
    denominator = math.sqrt(bcX ** 2 + bcY ** 2)
    return safe_div(numerator, denominator)


def calculate_foot_distance(A, B, C, D):
    if not all(is_valid_point(p) for p in [A, B, C, D]):
        return np.nan

    def get_foot(point):
        Bx, By = B
        Cx, Cy = C
        Px, Py = point
        vectorBCx = Cx - Bx
        vectorBCy = Cy - By
        vectorBPx = Px - Bx
        vectorBPy = Py - By
        denominator = vectorBCx ** 2 + vectorBCy ** 2
        if denominator == 0:
            return Bx, By
        t = (vectorBPx * vectorBCx + vectorBPy * vectorBCy) / denominator
        return Bx + t * vectorBCx, By + t * vectorBCy

    footA = get_foot(A)
    footD = get_foot(D)
    return math.sqrt((footA[0] - footD[0]) ** 2 + (footA[1] - footD[1]) ** 2)


def calculate_perpendicular_foot_distance(A, B, C, D):
    if not all(is_valid_point(p) for p in [A, B, C, D]):
        return np.nan
    Ax, Ay = A
    Bx, By = B
    Cx, Cy = C
    Dx, Dy = D
    abDx = Bx - Ax
    abDy = By - Ay
    denominator = abDx ** 2 + abDy ** 2
    if denominator == 0:
        return np.nan
    numerator = abDx * (Dx - Cx) + abDy * (Dy - Cy)
    t = numerator / denominator
    Fx = Dx - abDx * t
    Fy = Dy - abDy * t
    return math.sqrt((Fx - Cx) ** 2 + (Fy - Cy) ** 2)


# ========================
# 自然头位“前后/上下”带符号几何（前方=+x，下方=+y）
# ========================
def _perpendicular_foot(point, line_b, line_c):
    """点 point 在直线 line_b-line_c 上的垂足坐标。"""
    bx, by = line_b
    cx, cy = line_c
    px, py = point
    dx = cx - bx
    dy = cy - by
    denom = dx * dx + dy * dy
    if denom == 0:
        return bx, by
    t = ((px - bx) * dx + (py - by) * dy) / denom
    return bx + t * dx, by + t * dy


def signed_perp_distance_anterior(point, line_b, line_c):
    """点到直线的垂直距离，按自然头位前后定号：点在直线前方(+x侧)为正，后方为负。"""
    if not all(is_valid_point(p) for p in [point, line_b, line_c]):
        return np.nan
    foot = _perpendicular_foot(point, line_b, line_c)
    magnitude = math.hypot(point[0] - foot[0], point[1] - foot[1])
    dx = point[0] - foot[0]
    if dx == 0:
        return magnitude
    return magnitude if dx * ANTERIOR_X > 0 else -magnitude


def signed_angle_anb(a_pt, n_pt, b_pt):
    """ANB 角的带符号值：A 在 B 前方(+x侧)为正(III类负)，后方为负。"""
    magnitude = calculate_angle_at_center(a_pt, n_pt, b_pt)
    if pd.isna(magnitude) or not (is_valid_point(a_pt) and is_valid_point(b_pt)):
        return magnitude
    dx = a_pt[0] - b_pt[0]
    if dx == 0:
        return magnitude
    return magnitude if dx * ANTERIOR_X > 0 else -magnitude


def anb_from_sna_snb(s_pt, n_pt, a_pt, b_pt):
    """ANB = SNA - SNB。A 偏前(II类)为正，A 偏后(III类)为负。"""
    sna = calculate_angle_at_center(s_pt, n_pt, a_pt)
    snb = calculate_angle_at_center(s_pt, n_pt, b_pt)
    if pd.isna(sna) or pd.isna(snb):
        return np.nan
    return sna - snb


def signed_angle_sn_fh(s_pt, n_pt, po_pt, or_pt):
    """SN-FH：SN 平面与 FH 平面的前交角。
    SN 平面相对 FH 向上倾斜（两平面在后方相交）为正；向下倾斜（前方相交）为负。"""
    magnitude = calculate_angle_between_vectors(s_pt, n_pt, po_pt, or_pt)
    if pd.isna(magnitude) or not all(is_valid_point(p) for p in [s_pt, n_pt, po_pt, or_pt]):
        return magnitude
    # 两条由后向前(+x)的方向向量：SN 用 S->N，FH 用 Po->Or
    v_sn = (n_pt[0] - s_pt[0], n_pt[1] - s_pt[1])
    v_fh = (or_pt[0] - po_pt[0], or_pt[1] - po_pt[1])
    cross = v_fh[0] * v_sn[1] - v_fh[1] * v_sn[0]  # v_fh × v_sn
    if cross == 0:
        return magnitude
    # 图像坐标 y 向下：SN 相对 FH 向上倾斜时 cross<0，定为正值
    return magnitude if cross * ANTERIOR_X < 0 else -magnitude


def signed_horizontal_distance(point, ref_point, positive_when_front=True):
    """点到“过 ref_point 的竖直参考线”的水平距离。
    positive_when_front=True：point 在 ref 前方(+x侧)为正；False：point 在 ref 后方为正。"""
    if not is_valid_point(point) or not is_valid_point(ref_point):
        return np.nan
    front_signed = (point[0] - ref_point[0]) * ANTERIOR_X
    return front_signed if positive_when_front else -front_signed


def signed_overbite(u1p, l1p, occ_a, occ_b):
    """覆牙合(overbite)：沿合平面(occ_a-occ_b)的法线方向，U1-p 相对 L1-p 偏下方(+y)为正。"""
    if not all(is_valid_point(p) for p in [u1p, l1p, occ_a, occ_b]):
        return np.nan
    ux = occ_b[0] - occ_a[0]
    uy = occ_b[1] - occ_a[1]
    if ux == 0 and uy == 0:
        return np.nan
    # 合平面法线，取指向下方(+y)的方向
    perp_x, perp_y = -uy, ux
    if perp_y < 0:
        perp_x, perp_y = uy, -ux
    pnorm = math.hypot(perp_x, perp_y)
    return ((u1p[0] - l1p[0]) * perp_x + (u1p[1] - l1p[1]) * perp_y) / pnorm


def signed_overjet(u1p, l1p, occ_a, occ_b):
    """覆盖(overjet)：沿合平面(occ_a-occ_b)方向，U1-p 相对 L1-p 偏前方(+x侧)为正。"""
    if not all(is_valid_point(p) for p in [u1p, l1p, occ_a, occ_b]):
        return np.nan
    ux = occ_b[0] - occ_a[0]
    uy = occ_b[1] - occ_a[1]
    norm = math.hypot(ux, uy)
    if norm == 0:
        return np.nan
    # 合平面方向，取 x 分量指向前方(+x)的方向
    if ux * ANTERIOR_X < 0:
        ux, uy = -ux, -uy
    return ((u1p[0] - l1p[0]) * ux + (u1p[1] - l1p[1]) * uy) / norm


def rotate_points(points: Dict[str, Tuple[float, float]], angle_deg: float):
    """
    严格使用 rotation.csv 的角度旋转点。
    原参考脚本围绕图像中心旋转；本任务只有坐标 CSV，无图像尺寸。
    对本脚本计算的角度、距离、自然头位 x/y 差值而言，旋转中心差异只产生统一平移，差值不受影响。
    """
    if angle_deg == 0:
        return points
    angle_rad = math.radians(angle_deg)
    cos_a = math.cos(angle_rad)
    sin_a = math.sin(angle_rad)
    rotated = {}
    for key, point in points.items():
        if not is_valid_point(point):
            rotated[key] = point
            continue
        x, y = point
        rotated[key] = (x * cos_a - y * sin_a, x * sin_a + y * cos_a)
    return rotated


def rotate_points_around_center(points: Dict[str, Tuple[float, float]], angle_deg: float, center: Tuple[float, float]):
    """
    围绕图像中心顺时针旋转标志点 angle_deg 度。
    rotation.csv 中的角度约定为顺时针；该公式与 cv2.getRotationMatrix2D(center, -angle, 1.0) 一致。
    """
    if angle_deg == 0:
        return dict(points)
    cx, cy = center
    angle_rad = math.radians(angle_deg)
    cos_a = math.cos(angle_rad)
    sin_a = math.sin(angle_rad)
    rotated = {}
    for key, point in points.items():
        if not is_valid_point(point):
            rotated[key] = point
            continue
        x, y = point
        x_shifted = x - cx
        y_shifted = y - cy
        x_rot = x_shifted * cos_a - y_shifted * sin_a
        y_rot = x_shifted * sin_a + y_shifted * cos_a
        rotated[key] = (x_rot + cx, y_rot + cy)
    return rotated


def metric_value_dict(points_dict: Dict[str, Tuple[float, float]]) -> Dict[str, float]:
    r1 = points_dict.get("R1")
    r2 = points_dict.get("R2")
    scale = calculate_distance(r1, r2) / 10 if is_valid_point(r1) and is_valid_point(r2) else np.nan

    def p(name):
        return points_dict.get(name) if is_valid_point(points_dict.get(name)) else None

    s_point = p("S")
    n_point = p("N")
    ba_point = p("Ba")
    po_point = p("FH-a")
    or_point = p("FH-p")
    a_point = p("A")
    pns = p("PP-p")
    ans = p("PP-a")
    co = p("Co")
    rlc = p("Rlar-p")
    go = p("Go")
    me = p("MP-p")
    mp_p = p("MP-a")
    b_point = p("B")
    pog = p("Po")
    gn = p("GN")
    u1p = p("U1-p")
    u1a = p("U1-a")
    u1ta = p("U1-ta")
    u1c = p("U1C")
    l1c = p("L1C")
    l1p = p("L1-p")
    l1a = p("L1-a")
    u6a = p("U6-a")
    l6a = p("L6-a")
    u6p = p("U6-p")
    l6p = p("L6-p")
    l4 = p("L4")
    sn = p("Sn")
    glabella = p("G")
    ul = p("UL")
    ll = p("LL")
    ms = p("Ms")
    s_pog = p("Pos")
    stoms = p("Stoms")
    stomi = p("Stomi")

    ul6a_m = [(u6a[0] + l6a[0]) / 2, (u6a[1] + l6a[1]) / 2] if is_valid_point(u6a) and is_valid_point(l6a) else None
    ul6p_m = [(u6p[0] + l6p[0]) / 2, (u6p[1] + l6p[1]) / 2] if is_valid_point(u6p) and is_valid_point(l6p) else None
    ul1p_m = [(u1p[0] + l1p[0]) / 2, (u1p[1] + l1p[1]) / 2] if is_valid_point(u1p) and is_valid_point(l1p) else None

    return {
        "SN-FH(°)": signed_angle_sn_fh(s_point, n_point, po_point, or_point),
        "BaN-FH(°)": calculate_angle_between_vectors(ba_point, n_point, po_point, or_point),
        "S-N(mm)": safe_div(calculate_distance(s_point, n_point), scale),
        "S-Ba(mm)": safe_div(calculate_distance(s_point, ba_point), scale),
        "SNA(°)": calculate_angle_at_center(s_point, n_point, a_point),
        "PP-FH(°)": calculate_angle_between_vectors(po_point, or_point, pns, ans),
        "ANS-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(ans, n_point), scale),
        "A-N(自然)(mm)": safe_div(signed_horizontal_distance(a_point, n_point, positive_when_front=True), scale),
        "A-NP(mm)": safe_div(calculate_signed_vertical_distance(n_point, po_point, or_point, a_point), scale),
        "PNS-A(mm)": safe_div(calculate_foot_distance(a_point, pns, ans, pns), scale),
        "Co-A(mm)": safe_div(calculate_distance(co, a_point), scale),
        "Rlc-Go-Me(°)": calculate_angle_at_center(rlc, go, me),
        "Rlc-Go-N(°)": calculate_angle_at_center(rlc, go, n_point),
        "N-Go-Me(°)": calculate_angle_at_center(n_point, go, me),
        "N-S-Rlc(°)": calculate_angle_at_center(n_point, s_point, rlc),
        "S-Rlc-Go(°)": calculate_angle_at_center(s_point, rlc, go),
        "SNB(°)": calculate_angle_at_center(s_point, n_point, b_point),
        "Pog-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, pog), scale),
        "Pog-NB(mm)": safe_div(signed_perp_distance_anterior(pog, n_point, b_point), scale),
        "Pog-B(自然)(mm)": safe_div(calculate_natural_pos_level_distance(b_point, pog), scale),
        "B-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, b_point), scale),
        "B-NP(mm)": safe_div(calculate_signed_vertical_distance(n_point, po_point, or_point, b_point), scale),
        "Pog-NP(mm)": safe_div(calculate_signed_vertical_distance(n_point, po_point, or_point, pog), scale),
        "MP-FH(°)": calculate_angle_between_vectors(po_point, or_point, go, gn),
        "MP-SN(°)": calculate_angle_between_vectors(s_point, n_point, go, gn),
        "N-S-Gn(°)": calculate_angle_at_center(n_point, s_point, gn),
        "SGn-BaN(°)": calculate_angle_between_vectors(s_point, gn, n_point, ba_point),
        "SGn-FH(°)": calculate_angle_between_vectors(po_point, or_point, s_point, gn),
        "FH-NPog(°)": calculate_angle_between_vectors(po_point, or_point, pog, n_point),
        "Go-Pog(mm)": safe_div(calculate_distance(go, pog), scale),
        "Co-Go(mm)": safe_div(calculate_distance(go, co), scale),
        "Co-Gn(mm)": safe_div(calculate_distance(gn, co), scale),
        "ANB(°)": anb_from_sna_snb(s_point, n_point, a_point, b_point),
        "PP-GoGn(°)": calculate_angle_between_vectors(go, gn, pns, ans),
        "S-Go/N-Me(%)": safe_div(calculate_distance(s_point, go), calculate_distance(n_point, me)) * 100,
        "ANS-Me/N-Me(%)": safe_div(calculate_foot_to_c_distance(ans, n_point, me), calculate_distance(n_point, me)) * 100,
        "N-ANS(mm)": safe_div(calculate_foot_to_c_distance(ans, me, n_point), scale),
        "U1-A(mm)": safe_div(calculate_signed_vertical_distance(a_point, po_point, or_point, u1p), scale),
        "U1-NA(mm)": safe_div(signed_perp_distance_anterior(u1p, n_point, a_point), scale),
        "U1-PP(mm)": safe_div(calculate_perpendicular_distance_to_line(u1p, ans, pns), scale),
        "U1-APog(mm)": safe_div(signed_perp_distance_anterior(u1p, a_point, pog), scale),
        "U1p-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, u1p), scale),
        "U1p-Pog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(pog, u1p), scale),
        "U1ta-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, u1ta), scale),
        "U1ta-Pog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(pog, u1ta), scale),
        "U1c-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, u1c), scale),
        "U1c-Pog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(pog, u1c), scale),
        "L1-Pog(mm)": safe_div(calculate_signed_vertical_distance(pog, po_point, or_point, l1p), scale),
        "L1-NB(mm)": safe_div(signed_perp_distance_anterior(l1p, n_point, b_point), scale),
        "L1-MP(mm)": safe_div(calculate_perpendicular_distance_to_line(l1p, me, mp_p), scale),
        "L1-APog(mm)": safe_div(signed_perp_distance_anterior(l1p, a_point, pog), scale),
        "L1p-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, l1p), scale),
        "L1p-Pog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(pog, l1p), scale),
        "L1c-N(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, l1c), scale),
        "L1c-Pog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(pog, l1c), scale),
        "Holdaway ratio": safe_div(calculate_perpendicular_distance_to_line(pog, n_point, b_point), calculate_perpendicular_distance_to_line(l1p, n_point, b_point)),
        "U1-SN(°)": calculate_angle_between_vectors(s_point, n_point, u1p, u1a),
        "U1-NA(°)": calculate_angle_between_vectors(a_point, n_point, u1p, u1a),
        "L1-NB(°)": calculate_angle_between_vectors(l1a, l1p, b_point, n_point),
        "L1-MP(°)": calculate_angle_between_vectors(l1a, l1p, gn, go),
        "L1Pog-OP(°)": calculate_angle_at_center(pog, l1p, l6a),
        "U1-L1(°)": calculate_angle_between_vectors(u1p, u1a, l1p, l1a),
        "U6-PP(mm)": safe_div(calculate_perpendicular_distance_to_line(u6a, ans, pns), scale),
        "L6-MP(mm)": safe_div(calculate_perpendicular_distance_to_line(l6a, me, mp_p), scale),
        "overbite(mm)": safe_div(signed_overbite(u1p, l1p, ul6a_m, ul1p_m), scale),
        "overjet(mm)": safe_div(signed_overjet(u1p, l1p, ul6a_m, ul1p_m), scale),
        "OPa-SN(°)": calculate_angle_between_vectors(s_point, n_point, ul6a_m, ul1p_m),
        "OPa-FH(°)": calculate_angle_between_vectors(po_point, or_point, ul6a_m, ul1p_m),
        "OPf-SN(°)": calculate_angle_between_vectors(s_point, n_point, ul6a_m, l4),
        "OPf-FH(°)": calculate_angle_between_vectors(po_point, or_point, ul6a_m, l4),
        "L1-OP(mm)": safe_div(calculate_perpendicular_distance_to_line(l1p, ul6a_m, l4), scale),
        "sSn-glabella(自然)(mm)": safe_div(calculate_natural_pos_level_distance(glabella, sn), scale),
        "UL-Sn(自然)(mm)": safe_div(signed_horizontal_distance(ul, sn, positive_when_front=True), scale),
        "LL-Sn(自然)(mm)": safe_div(signed_horizontal_distance(ll, sn, positive_when_front=True), scale),
        "Ms-Sn(自然)(mm)": safe_div(signed_horizontal_distance(ms, sn, positive_when_front=False), scale),
        "sPog-Sn(自然)(mm)": safe_div(signed_horizontal_distance(s_pog, sn, positive_when_front=False), scale),
        "Sn-Stoms(mm)": safe_div(calculate_distance(sn, stoms), scale),
        "LL-sPogUL(mm)": safe_div(calculate_signed_perpendicular_distance(ll, ul, s_pog), scale),
        "Gl-Sn-sPog(°)": calculate_angle_at_center(glabella, sn, s_pog),
        "ANS-Stoms(mm)": safe_div(calculate_distance(ans, stoms), scale),
        "Sn-A(自然)(mm)": safe_div(calculate_natural_pos_level_distance(a_point, sn), scale),
        "Sn-N h(自然)(mm)": safe_div(calculate_natural_pos_level_distance(n_point, sn), scale),
        "Sn-N v(自然)(mm)": safe_div(calculate_natural_pos_vertical_distance(n_point, sn), scale),
        "Sn-ANS h(自然)(mm)": safe_div(calculate_natural_pos_level_distance(ans, sn), scale),
        "Sn-ANS v(自然)(mm)": safe_div(calculate_natural_pos_vertical_distance(ans, sn), scale),
        "U1p-Sn(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1p, sn), scale),
        "U1p-sPog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1p, s_pog), scale),
        "U1p-UL(mm)": safe_div(calculate_natural_pos_level_distance(u1p, ul), scale),
        "U1ta-Sn(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1ta, sn), scale),
        "U1ta-sPog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1ta, s_pog), scale),
        "U1ta-UL(mm)": safe_div(calculate_natural_pos_level_distance(u1ta, ul), scale),
        "U1c-Sn(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1c, sn), scale),
        "U1c-sPog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(u1c, s_pog), scale),
        "U1c-UL(mm)": safe_div(calculate_natural_pos_level_distance(u1c, ul), scale),
        "L1p-Sn(自然)(mm)": safe_div(calculate_natural_pos_level_distance(l1p, sn), scale),
        "L1p-sPog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(l1p, s_pog), scale),
        "L1p-UL(mm)": safe_div(calculate_natural_pos_level_distance(l1p, ul), scale),
        "L1c-Sn(自然)(mm)": safe_div(calculate_natural_pos_level_distance(l1c, sn), scale),
        "L1c-sPog(自然)(mm)": safe_div(calculate_natural_pos_level_distance(l1c, s_pog), scale),
        "L1c-UL(mm)": safe_div(calculate_natural_pos_level_distance(l1c, ul), scale),
        "Stomi-U1p(mm)": safe_div(calculate_natural_pos_vertical_distance(stomi, u1p), scale),
        "Stoms-U1p(mm)": safe_div(calculate_natural_pos_vertical_distance(stoms, u1p), scale),
        "Stomi-L1p(mm)": safe_div(calculate_natural_pos_vertical_distance(stomi, l1p), scale),
        "Stoms-L1p(mm)": safe_div(calculate_natural_pos_vertical_distance(stoms, l1p), scale),
    }


# ========================
# 数据读取与构造
# ========================
def patient_name_from_image(image_name: str) -> str:
    base = os.path.splitext(normalize_text(image_name))[0]
    if base.endswith("-1") or base.endswith("-2"):
        return base[:-2]
    return base


def image_suffix_from_name(image_name: str) -> str:
    base = os.path.splitext(normalize_text(image_name))[0]
    if base.endswith("-1") or base.endswith("-2"):
        return base[-2:]
    return "-1"


def find_image_path(patient_name: str, suffix: str):
    patient_dir = os.path.join(IMAGE_DIR, patient_name)
    if os.path.isdir(patient_dir):
        path = os.path.join(patient_dir, f"{patient_name}{suffix}{IMAGE_EXT}")
        if os.path.exists(path):
            return path
    direct_path = os.path.join(IMAGE_DIR, f"{patient_name}{suffix}{IMAGE_EXT}")
    if os.path.exists(direct_path):
        return direct_path
    return None


def get_image_center(image_path: str):
    if image_path in _image_center_cache:
        return _image_center_cache[image_path]
    with Image.open(image_path) as img:
        w, h = img.size
    center = (w / 2.0, h / 2.0)
    _image_center_cache[image_path] = center
    return center


def rotate_image_around_center(image, angle_deg):
    """围绕图像中心顺时针旋转图像 angle_deg 度。"""
    import cv2
    if angle_deg == 0:
        return image.copy()
    h, w = image.shape[:2]
    rotation_matrix = cv2.getRotationMatrix2D((w / 2.0, h / 2.0), -angle_deg, 1.0)
    return cv2.warpAffine(
        image, rotation_matrix, (w, h),
        borderMode=cv2.BORDER_CONSTANT, borderValue=(255, 255, 255),
    )


def draw_landmarks_on_image(image, points):
    import cv2
    result = image.copy()
    for label, point in points.items():
        if not is_valid_point(point):
            continue
        x, y = int(round(point[0])), int(round(point[1]))
        cv2.circle(result, (x, y), VIS_POINT_RADIUS, VIS_POINT_COLOR, -1)
        cv2.circle(result, (x, y), VIS_POINT_RADIUS, (0, 0, 0), 1)
        label_pos = (x + VIS_POINT_RADIUS + 2, y - VIS_POINT_RADIUS)
        cv2.putText(result, label, label_pos, cv2.FONT_HERSHEY_SIMPLEX, VIS_FONT_SCALE, (0, 0, 0), 3)
        cv2.putText(result, label, label_pos, cv2.FONT_HERSHEY_SIMPLEX, VIS_FONT_SCALE, VIS_TEXT_COLOR, 1)
    return result


def save_landmark_visualization(image_path: str, image_name: str, angle: float, rotated_points: Dict[str, Tuple[float, float]]):
    import cv2
    patient_name = patient_name_from_image(image_name)
    suffix = image_suffix_from_name(image_name)
    image = cv2.imread(image_path)
    if image is None:
        print(f"[Warning] 无法读取图像，跳过可视化: {image_path}")
        return

    rotated_image = rotate_image_around_center(image, angle)
    result = draw_landmarks_on_image(rotated_image, rotated_points)

    stage = "pre" if suffix == "-1" else "post"
    info_text = f"{stage}: {image_name} | Rotation: {angle:.1f}deg"
    cv2.putText(result, info_text, (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 0), 2)
    cv2.putText(result, info_text, (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 1)

    patient_output_dir = os.path.join(VIS_OUTPUT_DIR, patient_name)
    os.makedirs(patient_output_dir, exist_ok=True)
    output_path = os.path.join(patient_output_dir, f"{patient_name}_{stage}.png")
    cv2.imwrite(output_path, result)


def load_rotation_map(path: str) -> Dict[str, float]:
    rotation_map = {}
    with open(path, "r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        header = next(reader, None)
        for line_no, row in enumerate(reader, start=2):
            if len(row) != 2:
                print(f"[Warning] rotation.csv 第 {line_no} 行字段数异常，已跳过: {row}")
                continue
            name = normalize_text(row[0])
            try:
                rotation_map[name] = float(row[1])
            except ValueError:
                print(f"[Warning] rotation.csv 第 {line_no} 行角度不是数字，已跳过: {row}")
    return rotation_map


def load_points_csv(path: str) -> Dict[str, Dict[str, Tuple[float, float]]]:
    df = pd.read_csv(path, encoding="utf-8-sig")
    data = {}
    for _, row in df.iterrows():
        image_name = normalize_text(row["image"])
        points = {}
        for col in df.columns:
            if not col.endswith("_X"):
                continue
            label = col[:-2]
            y_col = f"{label}_Y"
            if y_col not in df.columns:
                continue
            x = pd.to_numeric(row[col], errors="coerce")
            y = pd.to_numeric(row[y_col], errors="coerce")
            points[label] = (float(x), float(y)) if not (pd.isna(x) or pd.isna(y)) else (np.nan, np.nan)
        data[image_name] = points
    return data


def find_patient_info_columns(sheet):
    first_header = [normalize_text(cell.value) for cell in sheet[1]]
    second_header = [normalize_text(cell.value) for cell in sheet[2]]
    filled_first_header = []
    current_header = ""
    for header in first_header:
        if header:
            current_header = header
        filled_first_header.append(current_header)

    name_col = None
    batch_col = None
    for col_idx, main_header in enumerate(filled_first_header, start=1):
        if main_header == "姓名":
            name_col = col_idx
        elif main_header == "批次":
            batch_col = col_idx

    if name_col is None or batch_col is None:
        raise ValueError("患者基本信息表中未找到 姓名 或 批次 列")
    return name_col, batch_col


def load_patient_split(path: str) -> Dict[str, str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    name_col, batch_col = find_patient_info_columns(sheet)
    split_map = {}
    skipped_blank_batch = 0

    for row_idx in range(3, sheet.max_row + 1):
        name = normalize_text(sheet.cell(row=row_idx, column=name_col).value)
        batch = normalize_batch(sheet.cell(row=row_idx, column=batch_col).value)
        if not name:
            continue
        if batch == "":
            skipped_blank_batch += 1
            continue
        if batch in TRAIN_BATCHES:
            split_map[name] = "train"
        elif batch in TEST_BATCHES:
            split_map[name] = "test"
        else:
            print(f"[Warning] 患者 {name} 批次={batch} 不在训练/测试规则内，已跳过")

    workbook.close()
    print(f"[Info] 患者批次表：有效划分 {len(split_map)} 位，批次为空跳过 {skipped_blank_batch} 位")
    return split_map


def calculate_metrics_for_image(
    image_name: str,
    points_map: Dict[str, Dict[str, Tuple[float, float]]],
    rotation_map: Dict[str, float],
    visualize: bool = False,
) -> Optional[Dict[str, float]]:
    points = points_map.get(image_name)
    if points is None:
        return None
    angle = rotation_map.get(image_name)
    if angle is None:
        print(f"[Warning] 缺少旋转角度: {image_name}，已使用 0")
        angle = 0.0

    patient_name = patient_name_from_image(image_name)
    suffix = image_suffix_from_name(image_name)
    image_path = find_image_path(patient_name, suffix)
    if image_path is None:
        print(f"[Warning] 找不到原始图像，无法获取图像中心，回退围绕原点旋转: {image_name}")
        rotated_points = rotate_points(points, angle)
    else:
        center = get_image_center(image_path)
        rotated_points = rotate_points_around_center(points, angle, center)
        if visualize:
            save_landmark_visualization(image_path, image_name, angle, rotated_points)

    return metric_value_dict(rotated_points)


def build_dataset():
    pre_measurements = read_measurement_list(PRE_MEASUREMENTS_FILE)
    delta_measurements = read_measurement_list(DELTA_MEASUREMENTS_FILE)
    post_measurements = read_measurement_list(POST_MEASUREMENTS_FILE)

    if len(pre_measurements) != 50:
        print(f"[Warning] pre_measurements.txt 当前读取到 {len(pre_measurements)} 个指标，不是描述中的 50 个")
    if len(delta_measurements) != 11:
        print(f"[Warning] post-pre_measurements.txt 当前读取到 {len(delta_measurements)} 个指标，不是描述中的 11 个")
    if len(post_measurements) != 6:
        print(f"[Warning] post_measurements.txt 当前读取到 {len(post_measurements)} 个指标，不是描述中的 6 个")

    patient_split = load_patient_split(PATIENT_INFO_EXCEL)
    rotation_map = load_rotation_map(ROTATION_CSV)
    pre_points = load_points_csv(PRE_POINTS_CSV)
    post_points = load_points_csv(POST_POINTS_CSV)

    if ENABLE_VISUALIZATION:
        os.makedirs(VIS_OUTPUT_DIR, exist_ok=True)
        print(f"[Info] 已启用可视化，标志点可视化将保存到: {VIS_OUTPUT_DIR}")

    rows = []
    skipped = []
    for patient_name, split in patient_split.items():
        pre_image = f"{patient_name}-1{IMAGE_EXT}"
        post_image = f"{patient_name}-2{IMAGE_EXT}"
        pre_metrics = calculate_metrics_for_image(pre_image, pre_points, rotation_map, visualize=ENABLE_VISUALIZATION)
        post_metrics = calculate_metrics_for_image(post_image, post_points, rotation_map, visualize=ENABLE_VISUALIZATION)

        if pre_metrics is None or post_metrics is None:
            skipped.append({"姓名": patient_name, "原因": "缺少 pre 或 post 标注点"})
            continue

        row = {"姓名": patient_name, "split": split}
        has_nan = False

        for metric in pre_measurements:
            value = pre_metrics.get(metric, np.nan)
            row[f"pre__{metric}"] = value
            has_nan = has_nan or pd.isna(value)

        for metric in delta_measurements:
            pre_value = pre_metrics.get(metric, np.nan)
            post_value = post_metrics.get(metric, np.nan)
            value = post_value - pre_value if not (pd.isna(pre_value) or pd.isna(post_value)) else np.nan
            row[f"delta_post_pre__{metric}"] = value
            has_nan = has_nan or pd.isna(value)

        for metric in post_measurements:
            value = post_metrics.get(metric, np.nan)
            row[f"target_post__{metric}"] = value
            has_nan = has_nan or pd.isna(value)

        if has_nan:
            skipped.append({"姓名": patient_name, "原因": "指标计算出现空值或无效值"})
            continue

        rows.append(row)

    all_df = pd.DataFrame(rows)
    train_df = all_df[all_df["split"] == "train"].copy()
    test_df = all_df[all_df["split"] == "test"].copy()

    feature_columns = [f"pre__{metric}" for metric in pre_measurements] + [f"delta_post_pre__{metric}" for metric in delta_measurements]
    target_columns = [f"target_post__{metric}" for metric in post_measurements]
    id_columns = ["姓名"]

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    # 按要求：一个任务只产出一个 all_dataset.csv（覆盖），训/测划分由后续训练代码基于此文件完成。
    all_df.to_csv(os.path.join(OUTPUT_DIR, "all_dataset.csv"), index=False, encoding="utf-8-sig")
    # 清理旧版遗留的子文件，保持输出目录整洁
    for legacy in ["train_features.csv", "train_targets.csv", "test_features.csv", "test_targets.csv", "skipped_patients.csv"]:
        legacy_path = os.path.join(OUTPUT_DIR, legacy)
        if os.path.exists(legacy_path):
            os.remove(legacy_path)

    print(f"[Success] 数据集已输出到: {OUTPUT_DIR}")
    print(f"[Info] 全部可用样本: {len(all_df)}")
    print(f"[Info] 训练集样本: {len(train_df)}")
    print(f"[Info] 测试集样本: {len(test_df)}")
    print(f"[Info] 输入特征数: {len(feature_columns)} = pre {len(pre_measurements)} + post-pre {len(delta_measurements)}")
    print(f"[Info] 输出标签数: {len(target_columns)}")
    print(f"[Info] 跳过患者: {len(skipped)}")


if __name__ == "__main__":
    build_dataset()
